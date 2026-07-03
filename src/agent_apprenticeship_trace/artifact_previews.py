from __future__ import annotations
import csv, json
from pathlib import Path
from typing import Any
from .io import sha256_file

TEXT_EXTS={'.txt','.md','.json','.jsonl','.csv','.py','.sh','.html','.xml','.xlsx'}

def _read_text(path: Path, limit: int) -> tuple[str, bool]:
    text=path.read_text(errors='replace')
    return text[:limit], len(text)>limit

def _preview_one(path: Path, package_ref: str, max_chars: int=4000, csv_rows: int=8) -> dict[str, Any]:
    size=path.stat().st_size if path.exists() else None
    out={'ref': package_ref, 'content_hash': ('sha256:' + sha256_file(path)) if path.exists() else None, 'size_bytes': size, 'media_type': 'unknown', 'preview': None, 'preview_truncated': False}
    ext=path.suffix.lower()
    if not path.exists():
        out.update({'parse_status':'missing'}); return out
    if ext not in TEXT_EXTS:
        out.update({'kind':'binary','media_type':'binary','parse_status':'metadata_only'}); return out
    out['media_type']='text'
    try:
        if ext == '.xlsx':
            try:
                from openpyxl import load_workbook  # type: ignore
            except Exception as exc:
                out.update({'kind':'xlsx','media_type':'data','parse_status':'openpyxl_unavailable','preview_error_message':str(exc)[:200]}); return out
            wb=load_workbook(path, data_only=False, read_only=False)
            sheets=[]; formulas=[]; important={'inputs','input','sensitivity','assumptions','summary','model','outputs'}
            for ws in wb.worksheets[:6]:
                rows=[]
                for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 20), max_col=min(ws.max_column or 1, 8), values_only=False):
                    vals=[]
                    for cell in row:
                        val=cell.value
                        vals.append(val)
                        if isinstance(val, str) and val.startswith('=') and len(formulas) < 40:
                            formulas.append({'sheet': ws.title, 'cell': cell.coordinate, 'formula': val})
                    rows.append(vals)
                sheets.append({'name': ws.title, 'max_row': ws.max_row, 'max_column': ws.max_column, 'headers': rows[0] if rows else [], 'first_rows': rows[1:]})
            out.update({'kind':'xlsx','media_type':'data','parse_status':'parsed_xlsx','sheet_names': wb.sheetnames, 'sheets': sheets, 'formulas_detected': bool(formulas), 'formulas': formulas, 'important_sheet_presence': {name: any(name in s.lower() for s in wb.sheetnames) for name in sorted(important)}, 'preview': json.dumps({'sheet_names': wb.sheetnames, 'sheets': sheets, 'formulas': formulas[:10]}, default=str)[:max_chars], 'preview_truncated': len(wb.sheetnames)>6 or len(formulas)>=40})
        elif ext == '.csv':
            text, trunc=_read_text(path, max_chars)
            with path.open(newline='', errors='replace') as f:
                reader=csv.reader(f); rows=[]
                for i,row in enumerate(reader):
                    rows.append(row)
                    if i >= csv_rows: break
            out.update({'kind':'csv','media_type':'data','parse_status':'parsed_csv','columns': rows[0] if rows else [], 'row_count': max(0, sum(1 for _ in path.open(errors='replace'))-1), 'first_rows': rows[1:], 'preview': text, 'preview_truncated': trunc})
        elif ext in {'.json','.jsonl'}:
            text, trunc=_read_text(path, max_chars)
            keys=[]; parse_status='parsed_json'
            if ext == '.json':
                try:
                    obj=json.loads(path.read_text(errors='replace'))
                    keys=sorted(obj.keys()) if isinstance(obj, dict) else []
                except Exception:
                    parse_status='json_parse_error'
            else:
                parse_status='jsonl_preview'
            out.update({'kind':'json' if ext == '.json' else 'jsonl','media_type':'data','parse_status':parse_status,'top_level_keys':keys,'preview':text,'preview_truncated':trunc})
        else:
            text, trunc=_read_text(path, max_chars)
            out.update({'kind':'text','parse_status':'text_preview','preview':text,'preview_truncated':trunc})
    except Exception as exc:
        out.update({'parse_status':'preview_error','preview_error_type':type(exc).__name__,'preview_error_message':str(exc)[:300]})
    return out

def build_artifact_previews(package_root: Path | None, refs: list[str], max_artifacts: int=12) -> dict[str, Any]:
    if package_root is None:
        return {'artifact_content_refs': [], 'artifact_content_previews': [], 'artifact_content_hashes': {}, 'artifact_content_preview_truncated': False, 'model_grading_basis': 'trace_only'}
    previews=[]
    for ref in list(dict.fromkeys(refs))[:max_artifacts]:
        p=package_root/ref
        if p.exists() and p.is_file():
            previews.append(_preview_one(p, ref))
    hashes={p['ref']:p.get('content_hash') for p in previews}
    truncated=any(bool(p.get('preview_truncated')) for p in previews)
    has_content=any(p.get('preview') for p in previews)
    return {'artifact_content_refs':[p['ref'] for p in previews], 'artifact_content_previews': previews, 'artifact_content_hashes': hashes, 'artifact_content_preview_truncated': truncated, 'model_grading_basis': 'artifact_content' if has_content and not truncated else ('artifact_preview' if previews else 'trace_only')}
